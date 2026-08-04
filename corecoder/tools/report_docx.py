"""Word 报告生成工具.

将 Agent 分析结果（Markdown 格式）生成专业的 Word (.docx) 报告文件。
使用 python-docx 实现，支持标题层级、粗体/斜体、代码块、表格等。
"""

import re
from io import BytesIO
from pathlib import Path

from .base import Tool

# Word 和 PPTX 报告统一输出目录
# ★ 修复: 统一输出到 web/backend/reports/，与下载端点保持一致
# 不再依赖 CWD，避免本地开发和服务器部署的路径差异
_REPORTS_DIR = (Path(__file__).resolve().parents[2] / "web" / "backend" / "reports").resolve()


class GenerateDocxReportTool(Tool):
    """将安全审计分析结果生成 Word 报告 (.docx 文件).

    接受 Markdown 格式的分析内容，生成格式化的 Word 文档，
    包含标题层级、粗体/斜体、代码块（等宽字体）、表格等专业排版。
    文件默认保存到 reports/ 目录下。
    适用于生成安全审计报告、漏洞分析报告、修复建议文档。
    """

    name = "generate_docx_report"
    description = (
        "将 Markdown 内容生成专业的 Word (.docx) 报告文件。"
        "支持标题（# ## ###）、粗体（**text**）、斜体（*text*）、"
        "行内代码（`code`）、代码块（```）、表格（|）、列表等完整 Markdown 语法。"
        "输出为 A4 格式、2.5cm 页边距、带页码的 Word 文档，"
        f"默认保存到 {_REPORTS_DIR / 'SECURITY_REPORT.docx'}。"
        "适用于：生成安全审计报告、创建漏洞分析文档、导出 Agent 分析结论。"
    )
    parameters = {
        "type": "object",
        "properties": {
            "content": {
                "type": "string",
                "description": "要生成报告的 Markdown 格式文本内容。通常为 Agent 分析完成后的综合结论。",
            },
            "output_path": {
                "type": "string",
                "description": (
                    "输出 .docx 文件的保存路径（绝对路径或相对路径）。"
                    f"默认保存到 {_REPORTS_DIR / 'SECURITY_REPORT.docx'}。"
                ),
            },
        },
        "required": ["content"],
    }

    def execute(self, content: str, output_path: str = "") -> str:
        """生成 Word 报告并保存到磁盘。

        Args:
            content: Markdown 格式的报告内容
            output_path: 可选输出路径，默认保存到 reports/SECURITY_REPORT.docx

        Returns:
            操作结果描述（含文件路径）
        """
        if not content or not content.strip():
            return "❌ 内容为空，无法生成报告。请先完成安全分析再生成报告。"

        # 解析输出路径
        if output_path:
            path = Path(output_path)
            if not path.is_absolute():
                path = Path.cwd() / path
        else:
            path = _REPORTS_DIR / "SECURITY_REPORT.docx"
        path.parent.mkdir(parents=True, exist_ok=True)

        try:
            docx_bytes = self._render_docx(content)
            path.write_bytes(docx_bytes)
        except Exception as e:
            import traceback
            traceback.print_exc()
            return f"❌ 报告生成失败：{e}"

        file_size_kb = len(docx_bytes) / 1024
        return (
            f"✅ Word 报告已生成！\n"
            f"📄 文件路径：{path}\n"
            f"📏 文件大小：{file_size_kb:.1f} KB\n"
            f"💡 可用 Microsoft Word / WPS / LibreOffice 打开查看。"
        )

    # ── 内部: Markdown → DOCX 渲染 ──────────────────────────────

    def _render_docx(self, markdown_content: str) -> bytes:
        """将 Markdown 内容转换为 Word 文档字节流."""
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        doc = Document()

        # ── 页面设置 ──
        section = doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)

        # ── 默认样式 ──
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        style.paragraph_format.line_spacing = 1.5

        # ── 解析 Markdown ──
        lines = markdown_content.split("\n")
        i = 0
        in_code_block = False
        code_lines = []

        while i < len(lines):
            line = lines[i]

            # 代码块
            if line.strip().startswith("```"):
                if in_code_block:
                    code_text = "\n".join(code_lines)
                    if code_text.strip():
                        para = doc.add_paragraph()
                        para.paragraph_format.space_before = Pt(6)
                        para.paragraph_format.space_after = Pt(6)
                        para.paragraph_format.left_indent = Cm(0.5)
                        run = para.add_run(code_text)
                        self._set_font(run, mono=True, size=9, color=(50, 50, 50))
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                i += 1
                continue

            if in_code_block:
                code_lines.append(line)
                i += 1
                continue

            # 空行
            if not line.strip():
                i += 1
                continue

            stripped = line.strip()

            # ── 标题 ──
            if stripped.startswith("# ") and not stripped.startswith("## "):
                self._add_heading_para(doc, stripped[2:], size=22, color=(30, 41, 59), center=True)
                # 标题装饰线
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(16)
                run = para.add_run("─" * 50)
                self._set_font(run, size=8, color=(180, 180, 180))

            elif stripped.startswith("## "):
                self._add_heading_para(doc, stripped[3:], size=16, color=(30, 41, 59))
            elif stripped.startswith("### "):
                self._add_heading_para(doc, stripped[4:], size=13, color=(51, 65, 85))
            elif stripped.startswith("#### "):
                self._add_heading_para(doc, stripped[5:], size=12, color=(71, 85, 105))

            # ── 水平线 ──
            elif stripped in ("---", "***", "___"):
                para = doc.add_paragraph()
                para.paragraph_format.space_before = Pt(8)
                para.paragraph_format.space_after = Pt(8)
                run = para.add_run("─" * 60)
                self._set_font(run, size=8, color=(200, 200, 200))

            # ── 表格 ──
            elif stripped.startswith("|") and stripped.endswith("|"):
                table_rows = []
                while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                    row_line = lines[i].strip()
                    if not re.match(r'^\|[\s\-:]+\|', row_line):
                        cells = [c.strip() for c in row_line[1:-1].split("|")]
                        table_rows.append(cells)
                    i += 1
                if table_rows:
                    self._add_wd_table(doc, table_rows)
                continue

            # ── 列表项 ──
            elif re.match(r'^[\s]*[\-\*\d+\.]\s', stripped):
                text = re.sub(r'^[\s]*[\-\*\d+\.]\s+', '', stripped)
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Cm(1.0)
                para.paragraph_format.space_after = Pt(3)
                self._add_inline_runs(para, text)

            # ── 引用块 ──
            elif stripped.startswith("> "):
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Cm(1.0)
                para.paragraph_format.space_after = Pt(4)
                run = para.add_run(stripped[2:])
                self._set_font(run, size=10, color=(100, 100, 100))

            # ── 普通段落 ──
            else:
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(6)
                self._add_inline_runs(para, stripped)
            i += 1

        # ── 页脚页码 ──
        for sec in doc.sections:
            footer = sec.footer
            footer.is_linked_to_previous = False
            fpara = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            fpara.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = fpara.add_run("— ")
            self._set_font(run, size=9, color=(150, 150, 150))
            fld_begin = OxmlElement("w:fldChar")
            fld_begin.set(qn("w:fldCharType"), "begin")
            run._r.append(fld_begin)
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = "PAGE"
            run._r.append(instr)
            fld_end = OxmlElement("w:fldChar")
            fld_end.set(qn("w:fldCharType"), "end")
            run._r.append(fld_end)
            run2 = fpara.add_run(" —")
            self._set_font(run2, size=9, color=(150, 150, 150))

        # ── 保存 ──
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.read()

    # ── 辅助方法 ──

    @staticmethod
    def _set_font(run, name="Calibri", size=11, bold=False, color=None, mono=False):
        from docx.shared import Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        western_font = "Consolas" if mono else name
        run.font.name = western_font
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)

        # 设置东亚字体为宋体，确保中文不依赖系统默认字体
        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:ascii'), western_font)
        rFonts.set(qn('w:hAnsi'), western_font)
        rFonts.set(qn('w:eastAsia'), '宋体')
        rFonts.set(qn('w:cs'), western_font)

    def _add_heading_para(self, doc, text, size=16, color=None, center=False):
        from docx.shared import Pt
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        para = doc.add_paragraph()
        run = para.add_run(text)
        self._set_font(run, size=size, bold=True, color=color)
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(4)
        if center:
            para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    def _add_inline_runs(self, para, text: str):
        """解析 inline markdown 并添加到段落."""
        from docx.shared import Pt
        pattern = re.compile(
            r'(\*\*(.+?)\*\*)|'        # **bold**
            r'(\*(.+?)\*)|'            # *italic*
            r'(`(.+?)`)'               # `code`
        )
        last_end = 0
        for match in pattern.finditer(text):
            prefix = text[last_end:match.start()]
            if prefix:
                run = para.add_run(prefix)
                self._set_font(run)
            if match.group(1):
                run = para.add_run(match.group(2))
                self._set_font(run, bold=True)
            elif match.group(3):
                run = para.add_run(match.group(4))
                run.italic = True
                self._set_font(run)
            elif match.group(5):
                run = para.add_run(match.group(6))
                self._set_font(run, mono=True, size=9.5)
            last_end = match.end()
        suffix = text[last_end:]
        if suffix:
            run = para.add_run(suffix)
            self._set_font(run)

    @staticmethod
    def _add_wd_table(doc, rows: list[list[str]]):
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        num_cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=num_cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row_data in enumerate(rows):
            for c_idx, cell_text in enumerate(row_data):
                if c_idx < num_cols:
                    cell = table.rows[r_idx].cells[c_idx]
                    cell.text = ""
                    cpara = cell.paragraphs[0]
                    crun = cpara.add_run(cell_text)
                    crun.font.name = "Calibri"
                    crun.font.size = Pt(10)
                    # 设置东亚字体为宋体
                    crPr = crun._r.get_or_add_rPr()
                    crFonts = crPr.find(qn('w:rFonts'))
                    if crFonts is None:
                        crFonts = OxmlElement('w:rFonts')
                        crPr.insert(0, crFonts)
                    crFonts.set(qn('w:ascii'), 'Calibri')
                    crFonts.set(qn('w:hAnsi'), 'Calibri')
                    crFonts.set(qn('w:eastAsia'), '宋体')
                    crFonts.set(qn('w:cs'), 'Calibri')
                    if r_idx == 0:
                        crun.bold = True
                        crun.font.color.rgb = RGBColor(255, 255, 255)
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "4472C4")
                        shading.set(qn("w:val"), "clear")
                        cell._tc.get_or_add_tcPr().append(shading)
                    elif r_idx % 2 == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "F2F6FC")
                        shading.set(qn("w:val"), "clear")
                        cell._tc.get_or_add_tcPr().append(shading)
        doc.add_paragraph()
