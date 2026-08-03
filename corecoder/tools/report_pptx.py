"""PPT 报告生成工具.

将安全审计发现（findings）生成 PowerPoint 执行摘要演示文稿。
使用 python-pptx 实现，包含标题页、统计摘要、漏洞详情列表和修复建议。
"""

import json
from pathlib import Path

from .base import Tool

# Word 和 PPTX 报告统一输出目录
_REPORTS_DIR = Path.cwd() / "reports"


class GeneratePptxReportTool(Tool):
    """将安全审计发现生成 PowerPoint 执行摘要 (.pptx 文件).

    接受审计发现列表或 JSON 文件路径，生成专业格式的 PPTX 演示文稿，
    包含标题页、统计摘要、漏洞详情和修复优先级建议。
    适用于向管理层汇报安全审计结果。
    """

    name = "generate_pptx_report"
    description = (
        "将安全审计发现（findings）生成 PowerPoint (.pptx) 执行摘要演示文稿。"
        "输入可以是 JSON 文件路径或直接传入 findings 数据。"
        "生成内容包括：标题页、统计摘要（按严重性分布）、"
        "Top 漏洞详情（严重性/CWE/文件/描述/修复建议）、修复优先级建议。"
        f"默认保存到 {_REPORTS_DIR / 'SECURITY_EXECUTIVE_SUMMARY.pptx'}。"
        "适用于：安全审计汇报、漏洞总结演示、管理层汇报。"
    )
    parameters = {
        "type": "object",
        "properties": {
            "findings_source": {
                "type": "string",
                "description": (
                    "审计发现的来源。可以是：\n"
                    "1) findings.json 文件的路径\n"
                    "2) 直接传入 JSON 字符串（findings 数组）\n"
                    "3) 留空则尝试读取 ./findings.json"
                ),
            },
            "output_path": {
                "type": "string",
                "description": (
                    "输出 .pptx 文件的保存路径。"
                    f"默认保存到 {_REPORTS_DIR / 'SECURITY_EXECUTIVE_SUMMARY.pptx'}。"
                ),
            },
            "title": {
                "type": "string",
                "description": "演示文稿标题（默认: 安全审计执行摘要）",
            },
        },
        "required": [],
    }

    def execute(
        self,
        findings_source: str = "",
        output_path: str = "",
        title: str = "",
    ) -> str:
        """生成 PPTX 执行摘要.

        Args:
            findings_source: 发现数据来源（JSON 文件路径或 JSON 字符串）
            output_path: 输出路径
            title: 演示文稿标题

        Returns:
            操作结果描述
        """
        # 解析 findings
        findings = self._load_findings(findings_source)
        if not findings:
            return (
                "❌ 未找到审计发现数据。\n"
                "请提供 findings.json 文件路径，或直接传入 findings JSON 数据。\n"
                "也可以先运行安全审计（c_review / insecure_defaults / injection_scanner）再生成报告。"
            )

        # 解析输出路径
        if output_path:
            path = Path(output_path)
            if not path.is_absolute():
                path = Path.cwd() / path
        else:
            path = _REPORTS_DIR / "SECURITY_EXECUTIVE_SUMMARY.pptx"
        path.parent.mkdir(parents=True, exist_ok=True)

        if not title:
            title = "安全审计执行摘要"

        try:
            self._render_pptx(findings, str(path), title)
        except Exception as e:
            import traceback
            traceback.print_exc()
            return f"❌ PPT 生成失败：{e}"

        file_size_kb = path.stat().st_size / 1024
        return (
            f"✅ PPT 执行摘要已生成！\n"
            f"📄 文件路径：{path}\n"
            f"📏 文件大小：{file_size_kb:.1f} KB\n"
            f"📊 包含 {len(findings)} 个漏洞发现\n"
            f"💡 可用 Microsoft PowerPoint / WPS / LibreOffice 打开放映。"
        )

    # ── 内部实现 ──────────────────────────────────────────────

    @staticmethod
    def _load_findings(source: str) -> list[dict]:
        """从多种来源加载 findings."""
        if not source:
            # 尝试默认路径
            default_path = Path.cwd() / "findings.json"
            if default_path.exists():
                return GeneratePptxReportTool._load_findings(str(default_path))
            return []

        # 尝试作为文件路径
        path = Path(source)
        if not path.is_absolute():
            path = Path.cwd() / path
        if path.exists() and path.suffix in (".json", ""):
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                if isinstance(data, list):
                    return data
                if isinstance(data, dict) and "findings" in data:
                    return data["findings"]
                return []
            except (json.JSONDecodeError, OSError):
                pass

        # 尝试作为 JSON 字符串
        try:
            data = json.loads(source)
            if isinstance(data, list):
                return data
            if isinstance(data, dict) and "findings" in data:
                return data["findings"]
        except json.JSONDecodeError:
            pass

        return []

    @staticmethod
    def _render_pptx(findings: list[dict], output_path: str, title: str):
        """渲染 PPTX 文件."""
        from pptx import Presentation
        from pptx.util import Inches, Pt, Emu
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.enum.shapes import MSO_SHAPE

        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9
        prs.slide_height = Inches(7.5)

        # 主题色
        DARK = RGBColor(30, 41, 59)
        PURPLE = RGBColor(99, 102, 241)
        WHITE = RGBColor(255, 255, 255)
        LIGHT_BG = RGBColor(248, 249, 255)
        GRAY = RGBColor(100, 116, 139)
        RED = RGBColor(239, 68, 68)
        ORANGE = RGBColor(249, 115, 22)
        YELLOW = RGBColor(234, 179, 8)
        GREEN = RGBColor(34, 197, 94)

        sev_order = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3, "INFORMATIONAL": 4}
        sorted_findings = sorted(findings, key=lambda f: sev_order.get(f.get("severity", "MEDIUM"), 5))

        # 统计
        by_sev = {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 0, "LOW": 0, "INFORMATIONAL": 0}
        for f in findings:
            sev = f.get("severity", "MEDIUM").upper()
            if sev in by_sev:
                by_sev[sev] += 1

        # ═══ Slide 1: 标题页 ═══
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        # 背景
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(30, 41, 59)

        # 装饰色块
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = PURPLE
        shape.line.fill.background()

        # 标题
        txBox = slide.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.3), Inches(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

        # 副标题
        txBox2 = slide.shapes.add_textbox(Inches(1.5), Inches(3.8), Inches(10.3), Inches(1.0))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = f"共发现 {len(findings)} 个安全漏洞 | CoreCoder 网络安全智能体"
        p2.font.size = Pt(18)
        p2.font.color.rgb = RGBColor(148, 163, 184)
        p2.alignment = PP_ALIGN.CENTER

        # ═══ Slide 2: 统计摘要 ═══
        slide2 = prs.slides.add_slide(prs.slide_layouts[6])
        # 顶部色条
        shape2 = slide2.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08)
        )
        shape2.fill.solid()
        shape2.fill.fore_color.rgb = PURPLE
        shape2.line.fill.background()

        # 标题
        txBox3 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf3 = txBox3.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = "📊 审计统计摘要"
        p3.font.size = Pt(32)
        p3.font.bold = True
        p3.font.color.rgb = DARK

        # 统计卡片
        sev_config = [
            ("CRITICAL", by_sev["CRITICAL"], RED, "🔴"),
            ("HIGH", by_sev["HIGH"], ORANGE, "🟠"),
            ("MEDIUM", by_sev["MEDIUM"], YELLOW, "🟡"),
            ("LOW", by_sev["LOW"], GREEN, "🟢"),
            ("INFO", by_sev["INFORMATIONAL"], GRAY, "🔵"),
        ]

        x_start = 0.5
        card_w = 2.2
        gap = 0.3
        for idx, (label, count, color, emoji) in enumerate(sev_config):
            left = Inches(x_start + idx * (card_w + gap))
            # 卡片背景
            card = slide2.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(card_w), Inches(2.0)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = LIGHT_BG
            card.line.color.rgb = RGBColor(226, 232, 240)

            # 数字
            txB = slide2.shapes.add_textbox(left, Inches(2.0), Inches(card_w), Inches(1.0))
            tBf = txB.text_frame
            pB = tBf.paragraphs[0]
            pB.text = str(count)
            pB.font.size = Pt(48)
            pB.font.bold = True
            pB.font.color.rgb = color
            pB.alignment = PP_ALIGN.CENTER

            # 标签
            txB2 = slide2.shapes.add_textbox(left, Inches(3.0), Inches(card_w), Inches(0.6))
            tBf2 = txB2.text_frame
            pB2 = tBf2.paragraphs[0]
            pB2.text = f"{emoji} {label}"
            pB2.font.size = Pt(16)
            pB2.font.color.rgb = GRAY
            pB2.alignment = PP_ALIGN.CENTER

        # 底部提示
        txBox4 = slide2.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(11.7), Inches(1.0))
        tf4 = txBox4.text_frame
        p4 = tf4.paragraphs[0]
        p4.text = f"🔍 共审计 {len(findings)} 个发现 | 涉及 {len(set(f.get('file', 'unknown') for f in findings))} 个文件"
        p4.font.size = Pt(14)
        p4.font.color.rgb = GRAY

        # ═══ Slide 3+: Top 漏洞详情 ═══
        top_findings = [f for f in sorted_findings if f.get("severity", "").upper() in ("CRITICAL", "HIGH")][:5]
        if not top_findings:
            top_findings = sorted_findings[:5]

        for f_idx, finding in enumerate(top_findings):
            slide_n = prs.slides.add_slide(prs.slide_layouts[6])
            sev = finding.get("severity", "MEDIUM")
            sev_color = {"CRITICAL": RED, "HIGH": ORANGE, "MEDIUM": YELLOW, "LOW": GREEN}.get(sev, GRAY)

            # 顶部色条
            shape_n = slide_n.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.06)
            )
            shape_n.fill.solid()
            shape_n.fill.fore_color.rgb = sev_color
            shape_n.line.fill.background()

            # 标题
            txT = slide_n.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.0))
            tTf = txT.text_frame
            pT = tTf.paragraphs[0]
            pT.text = f"🐛 {finding.get('id', f'FINDING-{f_idx+1}')}: {finding.get('title', '未命名漏洞')}"
            pT.font.size = Pt(24)
            pT.font.bold = True
            pT.font.color.rgb = DARK

            # 元数据标签行
            meta_items = [
                f"严重性: {sev}",
                f"CWE: {finding.get('cwe_id', 'N/A')}",
                f"CVSS: {finding.get('cvss_score', 'N/A')}",
                f"文件: {Path(finding.get('file', 'N/A')).name}",
                f"行号: {finding.get('line', 'N/A')}",
            ]
            txM = slide_n.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.5))
            tMf = txM.text_frame
            pM = tMf.paragraphs[0]
            pM.text = " | ".join(meta_items)
            pM.font.size = Pt(13)
            pM.font.color.rgb = GRAY

            # 描述
            desc = finding.get("description", "无详细描述")
            txD = slide_n.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.2))
            tDf = txD.text_frame
            tDf.word_wrap = True
            pD = tDf.paragraphs[0]
            pD.text = f"📋 描述\n{desc}"
            pD.font.size = Pt(14)
            pD.font.color.rgb = DARK
            pD.space_after = Pt(8)

            # 攻击场景
            attack = finding.get("attack_scenario", "")
            if attack:
                pD2 = tDf.add_paragraph()
                pD2.text = f"\n⚠️ 攻击场景\n{attack}"
                pD2.font.size = Pt(14)
                pD2.font.color.rgb = RED

            # 修复建议
            fix = finding.get("fix_suggestion", "")
            if fix:
                txF = slide_n.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(11.7), Inches(1.8))
                tFf = txF.text_frame
                tFf.word_wrap = True
                pF = tFf.paragraphs[0]
                pF.text = f"💡 修复建议\n{fix}"
                pF.font.size = Pt(14)
                pF.font.color.rgb = RGBColor(22, 163, 74)

        # ═══ 最后一张: 修复优先级建议 ═══
        slide_last = prs.slides.add_slide(prs.slide_layouts[6])
        shape_last = slide_last.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08)
        )
        shape_last.fill.solid()
        shape_last.fill.fore_color.rgb = PURPLE
        shape_last.line.fill.background()

        txL = slide_last.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.8))
        tLf = txL.text_frame
        pL = tLf.paragraphs[0]
        pL.text = "🔧 修复优先级建议"
        pL.font.size = Pt(32)
        pL.font.bold = True
        pL.font.color.rgb = DARK

        priorities = [
            ("1. 立即修复", f"所有 CRITICAL({by_sev['CRITICAL']}) 和 HIGH({by_sev['HIGH']}) 级别漏洞", RED),
            ("2. 本迭代修复", f"MEDIUM({by_sev['MEDIUM']}) 级别漏洞", ORANGE),
            ("3. 计划修复", f"LOW({by_sev['LOW']}) 级别漏洞，纳入后续迭代", YELLOW),
            ("4. 持续关注", f"INFORMATIONAL({by_sev['INFORMATIONAL']}) 级别作为安全基线参考", GRAY),
        ]

        for idx, (label, desc, color) in enumerate(priorities):
            y_pos = 1.8 + idx * 1.2
            # 序号圆圈
            circle = slide_last.shapes.add_shape(
                MSO_SHAPE.OVAL, Inches(1.5), Inches(y_pos), Inches(0.5), Inches(0.5)
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = color
            circle.line.fill.background()

            txP = slide_last.shapes.add_textbox(Inches(2.3), Inches(y_pos - 0.05), Inches(10.0), Inches(0.8))
            tPf = txP.text_frame
            pP = tPf.paragraphs[0]
            pP.text = label
            pP.font.size = Pt(20)
            pP.font.bold = True
            pP.font.color.rgb = DARK
            pP2 = tPf.add_paragraph()
            pP2.text = desc
            pP2.font.size = Pt(14)
            pP2.font.color.rgb = GRAY

        # 页脚
        txFoot = slide_last.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5))
        tFootf = txFoot.text_frame
        pFoot = tFootf.paragraphs[0]
        pFoot.text = "本报告由 CoreCoder 网络安全智能体自动生成 | 仅供内部安全审计使用"
        pFoot.font.size = Pt(10)
        pFoot.font.color.rgb = GRAY
        pFoot.alignment = PP_ALIGN.CENTER

        prs.save(output_path)
